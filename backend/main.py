import os
import io
import requests
from typing import List
from bs4 import BeautifulSoup
from docx import Document
import PyPDF2

from fastapi import FastAPI, HTTPException, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse, FileResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

try:
    from backend.summarizer import summarizer_instance
except ImportError:
    from summarizer import summarizer_instance

app = FastAPI(title="Text Summarization API", description="API for summarizing text using NLP models")

# Configure CORS for frontend access
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

class SummarizeRequest(BaseModel):
    text: str
    max_length: int = 130
    min_length: int = 30
    style: str = "abstractive"
    bullet_count: int = 3

class ScrapeRequest(BaseModel):
    url: str

class ExportRequest(BaseModel):
    summary: str
    keywords: List[str]
    sentiment: str

class SummarizeResponse(BaseModel):
    summary: str
    keywords: List[str] = []
    sentiment: str = "Neutral"

@app.get("/healthz", tags=["health"])
async def health():
    return {"status": "ok"}

@app.post("/summarize", response_model=SummarizeResponse)
async def summarize_text(request: SummarizeRequest):
    if not request.text:
        raise HTTPException(status_code=400, detail="Text cannot be empty")
    
    if request.style == "extractive":
        summary = summarizer_instance.extractive_summarize(
            text=request.text,
            sentences_count=request.bullet_count
        )
    else:
        summary = summarizer_instance.summarize(
            text=request.text,
            max_length=request.max_length,
            min_length=request.min_length
        )
    
    keywords = summarizer_instance.extract_keywords(request.text, num_keywords=5)
    sentiment = summarizer_instance.analyze_sentiment(request.text)
    
    return SummarizeResponse(summary=summary, keywords=keywords, sentiment=sentiment)

@app.post("/upload")
async def upload_file(file: UploadFile = File(...)):
    try:
        content = await file.read()
        text = ""
        
        if file.filename.endswith(".pdf"):
            pdf_reader = PyPDF2.PdfReader(io.BytesIO(content))
            for page in pdf_reader.pages:
                text += page.extract_text() + "\n"
        else:
            text = content.decode("utf-8")
            
        return {"text": text}
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Error processing file: {str(e)}")

@app.post("/scrape")
async def scrape_url(request: ScrapeRequest):
    try:
        headers = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)'}
        response = requests.get(request.url, headers=headers, timeout=10)
        response.raise_for_status()
        
        soup = BeautifulSoup(response.content, "html.parser")
        
        paragraphs = soup.find_all('p')
        text = "\n".join([p.get_text() for p in paragraphs])
        
        if len(text.strip()) < 50:
            text = soup.get_text(separator='\n', strip=True)
            
        return {"text": text}
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to scrape URL: {str(e)}")

@app.post("/export")
async def export_document(request: ExportRequest):
    try:
        doc = Document()
        doc.add_heading('AI Text Summary', 0)

        doc.add_heading('Sentiment', level=1)
        doc.add_paragraph(request.sentiment)

        doc.add_heading('Keywords', level=1)
        doc.add_paragraph(', '.join(request.keywords))

        doc.add_heading('Summary', level=1)
        doc.add_paragraph(request.summary)

        file_stream = io.BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)

        return StreamingResponse(
            file_stream, 
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": "attachment; filename=Summary.docx"}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to generate document: {str(e)}")

# Frontend static serving
_frontend_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "frontend"))
if os.path.exists(_frontend_dir):
    @app.get("/", include_in_schema=False)
    async def serve_index():
        return FileResponse(os.path.join(_frontend_dir, "index.html"))

    @app.get("/{file_name:path}", include_in_schema=False)
    async def serve_static(file_name: str):
        file_path = os.path.join(_frontend_dir, file_name)
        if os.path.isfile(file_path):
            return FileResponse(file_path)
        raise HTTPException(status_code=404, detail="Not found")

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
